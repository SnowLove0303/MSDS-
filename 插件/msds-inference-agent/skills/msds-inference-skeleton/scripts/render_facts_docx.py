#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the raw-input facts package as a clearly labelled Word report."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GRAY = "666666"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], BLUE)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if index == 1 and value in {"原始输入未提供", "需检索", "需人工审核"}:
                shade(cells[index], "FFF2CC")
    document.add_paragraph()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = f"Heading {min(level, 3)}"
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--facts", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    facts = json.loads(Path(args.facts).read_text(encoding="utf-8"))

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PEA-4139 原始输入事实\n数据库标准 17 节骨架映射报告")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Microsoft YaHei"
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subrun = subtitle.add_run("事实包 v0.2｜数据库标准 S9=37 项、S11=10 项｜非正式 MSDS｜不作为合规文件直接使用")
    subrun.font.size = Pt(10)
    subrun.font.color.rgb = RGBColor.from_string(GRAY)

    add_heading(document, "一、文档性质与来源", 1)
    document.add_paragraph("本报告按数据库 msds_standard.db 的 schema_field 标准展开 S0+S1–S16；其中 S9 固定 37 项、S11 固定 10 项。报告只呈现用户输入的原始事实、确定性派生字段、标准骨架要求和证据状态。原始值不因外部搜索而被改写；“原始输入未提供”“需检索”“需人工审核”均为事实包状态，不是正式 MSDS 文案。")
    add_table(document, ["项目", "值"], [
        ["源文件", facts["source"]["path"]],
        ["SHA-256", facts["source"]["sha256"]],
        ["文件大小", str(facts["source"]["size"])],
        ["工作表", "、".join(facts["source"]["sheets"])],
        ["骨架版本", facts["skeleton"]["version"]],
        ["骨架节数", str(facts["skeleton"]["section_count"])],
        ["正式 MSDS", "否"],
    ])

    add_heading(document, "二、标准与法规证据（与事实分离）", 1)
    evidence_rows = []
    for item in facts.get("regulatory_evidence", []):
        evidence_rows.append([item.get("id", ""), item.get("title", ""), item.get("status", ""), item.get("source_url", ""), item.get("retrieved_at", "")])
    add_table(document, ["ID", "标题", "状态", "来源", "读取日期"], evidence_rows or [["—", "尚未附证据", "需检索", "", ""]])
    if facts["skeleton"].get("known_baseline_conflicts"):
        document.add_paragraph("骨架基线冲突：" + "；".join(facts["skeleton"]["known_baseline_conflicts"]), style="Intense Quote")

    add_heading(document, "三、输入异常与审核边界", 1)
    for anomaly in facts.get("input_anomalies", []):
        document.add_paragraph(anomaly, style="List Bullet")

    add_heading(document, "四、17 节标准骨架事实", 1)
    for section_id in facts["skeleton"]["sections"]:
        section_data = facts["sections"][section_id]
        add_heading(document, f"Section {section_id}｜{section_data['title']}", 2)
        document.add_paragraph(f"节状态：{section_data['status']}")
        rows = []
        for field in section_data.get("fields", []):
            source = field.get("source_fact")
            if field.get("status") == "source_fact" and source:
                value = source.get("raw_value", "")
                origin = f"原始输入 {source.get('source_cells', {}).get('value', '')}"
                label = f"{field.get('label', '')}（原标签：{source.get('raw_label', '')}）"
            else:
                value = field.get("value", "")
                origin = field.get("status", "")
                label = field.get("label", "")
            rows.append([label, value, field.get("status", ""), origin])
        if rows:
            add_table(document, ["标准字段/原始标签", "呈现值", "状态", "来源/说明"], rows)
        if section_id == "3":
            component_rows = []
            for component in section_data.get("components", []):
                component_rows.append([component.get("raw_name", ""), component.get("raw_cas", ""), component.get("raw_concentration", ""), "原始输入"])
            add_table(document, ["化学品名称（原文）", "CAS编号（原文）", "含量（原文）", "状态"], component_rows or [["—", "—", "—", "原始输入未提供"]])

    add_heading(document, "五、结论边界", 1)
    document.add_paragraph("本报告已按现有 S0+S1–S16 逻辑骨架展开，但仅 S1、S3、S9 含用户输入事实；其余节没有被自动编造。若要形成正式 MSDS，必须在独立授权流程中完成成分身份、GHS 分类、混合物判定、职业接触限值、毒理/生态、运输、法规适用性及跨节一致性审核，并重新生成正式模板文件。")

    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"[OK] Word report: {output}")


if __name__ == "__main__":
    main()
