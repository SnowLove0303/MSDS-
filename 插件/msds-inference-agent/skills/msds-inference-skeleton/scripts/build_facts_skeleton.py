#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a source-preserving database-standard S0+S1-S16 facts package from an MSDS workbook.

This initial implementation deliberately does not infer a formal MSDS. It copies
S1/S3/S9 values, records coordinates and anomalies, and materializes every field
in the database-canonical 17-section skeleton (S9=37 fields, S11=10 fields).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAIN_SECTION_RE = re.compile(r"^\s*(\d+)\s*[.．、](?!\s*\d)")
SEQ_RE = re.compile(r"^\s*(\d+(?:[.．]\d+)*)\s*(.*)$")
FIELD_ALIASES = {
    "嗅觉阀值": "嗅觉阈值",
    "pH值（1%水溶液）": "pH值",
}


def raw_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def normalized_label(value: Any) -> str:
    text = raw_text(value).replace("\xa0", " ").replace("　", " ").strip()
    text = re.sub(r"[：:，,。；;]+$", "", text)
    text = re.sub(r"\s+", "", text)
    match = SEQ_RE.match(text)
    label = match.group(2) if match else text
    return FIELD_ALIASES.get(label, label)


def parsed_seq(value: Any) -> str:
    text = raw_text(value).replace("．", ".")
    match = SEQ_RE.match(text)
    return match.group(1) if match else ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_skeleton(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def fact(section: str, label_cell: str, value_cell: str, label: Any, value: Any) -> dict[str, Any]:
    return {
        "section": section,
        "status": "source_fact",
        "origin": "source",
        "source_cells": {"label": label_cell, "value": value_cell},
        "raw_label": raw_text(label),
        "normalized_label": normalized_label(label),
        "sequence": parsed_seq(label),
        "raw_value": raw_text(value),
        "value": raw_text(value),
    }


def workbook_rows(workbook: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            values = [cell.value for cell in row]
            if not any(value is not None for value in values):
                continue
            rows.append({
                "sheet": worksheet.title,
                "row": row[0].row,
                "cells": {cell.column_letter: cell.value for cell in row if cell.value is not None},
            })
    return rows


def find_sections(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    blocks: dict[str, list[dict[str, Any]]] = {}
    current: str | None = None
    for item in rows:
        a = item["cells"].get("A")
        match = MAIN_SECTION_RE.match(raw_text(a))
        if match:
            current = match.group(1)
        if current is not None:
            blocks.setdefault(current, []).append(item)
    return blocks


def parse_s1(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for item in rows:
        label = item["cells"].get("A")
        value = item["cells"].get("B")
        if MAIN_SECTION_RE.match(raw_text(label)) or not raw_text(value):
            continue
        output.append(fact("1", f"{item['sheet']}!A{item['row']}", f"{item['sheet']}!B{item['row']}", label, value))
    return output


def parse_s3(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    product_type: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    anomalies: list[str] = []
    in_component_table = False
    header_seen = False
    for item in rows:
        cells = item["cells"]
        a, b, c = cells.get("A"), cells.get("B"), cells.get("C")
        label = normalized_label(a)
        if MAIN_SECTION_RE.match(raw_text(a)):
            continue
        if label == "产品类型" and raw_text(b):
            product_type.append(fact("3", f"{item['sheet']}!A{item['row']}", f"{item['sheet']}!B{item['row']}", a, b))
            continue
        if label == "成分":
            in_component_table = True
            continue
        if not in_component_table:
            continue
        if not header_seen:
            header_seen = True
            continue
        if not any(raw_text(value) for value in (a, b, c)):
            continue
        components.append({
            "status": "source_fact",
            "origin": "source",
            "source_cells": {"name": f"{item['sheet']}!A{item['row']}", "cas": f"{item['sheet']}!B{item['row']}", "concentration": f"{item['sheet']}!C{item['row']}"},
            "raw_name": raw_text(a),
            "raw_cas": raw_text(b),
            "raw_concentration": raw_text(c),
            "name": raw_text(a),
            "cas": raw_text(b),
            "concentration": raw_text(c),
        })
    if not header_seen and components:
        anomalies.append("S3 成分表未识别到表头；成分行仍按原始列读取。")
    return product_type, components, anomalies


def parse_s9(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    output: list[dict[str, Any]] = []
    anomalies: list[str] = []
    seen_sequences: dict[str, list[str]] = {}
    for item in rows:
        label = item["cells"].get("A")
        value = item["cells"].get("B")
        if MAIN_SECTION_RE.match(raw_text(label)) or not raw_text(value):
            continue
        current = fact("9", f"{item['sheet']}!A{item['row']}", f"{item['sheet']}!B{item['row']}", label, value)
        output.append(current)
        sequence = current["sequence"]
        if sequence:
            seen_sequences.setdefault(sequence, []).append(current["normalized_label"])
    for sequence, labels in seen_sequences.items():
        if len(labels) > 1:
            anomalies.append(f"S9 序号 {sequence} 对应多个原始字段：{'；'.join(labels)}；已保留原始标签和值。")
    return output, anomalies


def missing_field(label: str, status: str = "source_not_provided", note: str | None = None) -> dict[str, Any]:
    return {
        "status": status,
        "origin": "skeleton_requirement",
        "label": label,
        "value": "原始输入未提供" if status == "source_not_provided" else ("需检索" if status == "needs_search" else "需人工审核"),
        "note": note or "该值不是来源事实，不能直接写入正式 MSDS。",
    }


def make_section(section_id: str, spec: dict[str, Any], source_facts: list[dict[str, Any]], evidence: list[dict[str, Any]]) -> dict[str, Any]:
    by_label = {item["normalized_label"]: item for item in source_facts}
    fields: list[dict[str, Any]] = []
    for label in spec.get("fields", []):
        source = by_label.get(FIELD_ALIASES.get(label, label))
        if source is not None:
            fields.append({"status": "source_fact", "origin": "source", "label": label, "value": source["raw_value"], "source_fact": source})
        else:
            fields.append(missing_field(label))
    return {
        "section": section_id,
        "title": spec["title"],
        "status": "source_facts_available" if source_facts else "source_not_provided",
        "fields": fields,
        "evidence_refs": [item.get("id") for item in evidence if item.get("section") in (None, section_id)],
    }


def build(args: argparse.Namespace) -> dict[str, Any]:
    source = Path(args.input).resolve()
    skeleton_path = Path(args.skeleton).resolve()
    workbook = load_workbook(source, data_only=False, read_only=False)
    rows = workbook_rows(workbook)
    blocks = find_sections(rows)
    s1 = parse_s1(blocks.get("1", []))
    s3_product_type, components, s3_anomalies = parse_s3(blocks.get("3", []))
    s9, s9_anomalies = parse_s9(blocks.get("9", []))
    evidence: list[dict[str, Any]] = []
    if args.evidence:
        evidence = json.loads(Path(args.evidence).read_text(encoding="utf-8"))
        if isinstance(evidence, dict):
            evidence = evidence.get("evidence", [])

    skeleton = load_skeleton(skeleton_path)
    sections: dict[str, Any] = {}
    for section_id, spec in skeleton["sections"].items():
        if section_id == "0":
            model = ""
            basis = None
            for item in s1:
                if item["normalized_label"] == "中文名称":
                    match = re.search(r"[A-Za-z]{1,8}-?\d{2,}", item["raw_value"])
                    if match:
                        model = match.group(0)
                        basis = item
                        break
            fields = []
            for label in spec.get("fields", []):
                if label in ("产品名称", "产品型号") and model:
                    fields.append({"status": "derived_fact", "origin": "derived", "label": label, "value": model, "basis": basis["source_cells"], "note": "从 S1 中文名称中的型号模式派生；非原始单元格事实。"})
                else:
                    fields.append(missing_field(label))
            sections[section_id] = {"section": section_id, "title": spec["title"], "status": "derived_and_source_not_provided", "fields": fields, "evidence_refs": []}
        elif section_id == "1":
            sections[section_id] = make_section(section_id, spec, s1, evidence)
        elif section_id == "3":
            section = make_section(section_id, spec, s3_product_type, evidence)
            section["components"] = components
            section["status"] = "source_facts_available"
            sections[section_id] = section
        elif section_id == "9":
            sections[section_id] = make_section(section_id, spec, s9, evidence)
        else:
            sections[section_id] = make_section(section_id, spec, [], evidence)

    anomalies = s3_anomalies + s9_anomalies
    canonical_s9 = {
        label: f"9.{index}"
        for index, label in enumerate(skeleton["sections"]["9"].get("fields", []), start=1)
    }
    for source_item in s9:
        canonical_seq = canonical_s9.get(source_item["normalized_label"])
        if canonical_seq and source_item["sequence"] and canonical_seq != source_item["sequence"]:
            anomalies.append(
                f"S9 原始字段 {source_item['raw_label']} 使用序号 {source_item['sequence']}；"
                f"数据库标准字段 {source_item['normalized_label']} 序号为 {canonical_seq}；"
                "原始标签和值保持不变。"
            )
    if not any(item["normalized_label"] in {"供应商名称", "供应商地址", "电话", "传真"} for item in s1):
        anomalies.append("S1 未提供供应商信息字段；不得从现有模板或历史产物补写。")
    if components and any(item["raw_concentration"].lstrip().startswith((">", "<")) for item in components):
        anomalies.append("S3 至少一个成分含量使用不等式区间；本事实包不把它转换为精确百分比。")
    anomalies.append("本文件是原始输入事实包，不是正式 MSDS；缺失字段不代表法律上的‘无数据’或‘不适用’结论。")

    return {
        "schema_version": "msds-17-skeleton-facts-db-v0.2",
        "document_type": "raw_input_facts",
        "formal_msds": False,
        "source": {
            "path": str(source),
            "sha256": sha256(source),
            "size": source.stat().st_size,
            "sheets": workbook.sheetnames,
            "nonempty_row_count": len(rows),
        },
        "skeleton": {
            "path": str(skeleton_path),
            "version": skeleton.get("schema_version"),
            "section_count": len(sections),
            "sections": list(sections),
            "known_baseline_conflicts": skeleton.get("known_baseline_conflicts", []),
        },
        "regulatory_evidence": evidence,
        "input_anomalies": anomalies,
        "source_facts": {"s1": s1, "s3_product_type": s3_product_type, "s3_components": components, "s9": s9},
        "sections": sections,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Input .xlsx path")
    parser.add_argument("--out", required=True, help="Output JSON path")
    parser.add_argument("--evidence", help="Optional regulatory evidence JSON")
    parser.add_argument("--skeleton", default=str(Path(__file__).resolve().parents[1] / "references" / "standard_skeleton.json"))
    args = parser.parse_args()
    result = build(args)
    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] facts package: {output}")
    print(f"[INFO] sections={result['skeleton']['section_count']} source_sha256={result['source']['sha256']}")
    print(f"[INFO] anomalies={len(result['input_anomalies'])} evidence={len(result['regulatory_evidence'])}")


if __name__ == "__main__":
    main()
