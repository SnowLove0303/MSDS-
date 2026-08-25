# -*- coding: utf-8 -*-
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
from content_part1 import INTRO, DIAG_TABLE_HEADER, DIAG_TABLE_ROWS, BASIS
from content_part2 import SECTIONS_1_8
from content_part3 import SECTIONS_9_16, REFERENCES

doc = Document()

# 页面边距
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)

def set_font(run, name_cn='宋体', name_en='Times New Roman', size=10.5, bold=False, color=None):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def para(text, size=10.5, bold=False, cn='宋体', align=None, space_after=6, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, name_cn=cn, size=size, bold=bold, color=color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(text, level=1):
    sizes = {0:16, 1:14, 2:12, 3:11}
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, name_cn='黑体', size=sizes.get(level,11), bold=True, color=(0x1F,0x4E,0x79))
    p.paragraph_format.space_before = Pt(12 if level>0 else 6)
    p.paragraph_format.space_after = Pt(6)
    return p

# ===== 封面标题 =====
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(INTRO["title"]); set_font(r, name_cn='黑体', size=18, bold=True, color=(0x1F,0x4E,0x79))
p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("——依据第1、3、9部分推导其余各部分的编写方案"); set_font(r2, name_cn='楷体', size=13, color=(0x44,0x44,0x44))
doc.add_paragraph()

# 元信息
for line in INTRO["meta"].split("\n"):
    para(line, size=10, cn='宋体', space_after=3)

# ===== 一、概述 =====
heading("一、任务概述与原文诊断", 0)
for t in INTRO["overview"]:
    para(t)

heading("1.1 原文主要规范性问题汇总", 2)
tbl = doc.add_table(rows=1, cols=3)
tbl.style = 'Table Grid'
for i, h in enumerate(DIAG_TABLE_HEADER):
    c = tbl.rows[0].cells[i]
    c.text = ''
    r = c.paragraphs[0].add_run(h); set_font(r, name_cn='黑体', size=10, bold=True)
for row in DIAG_TABLE_ROWS:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = ''
        r = cells[i].paragraphs[0].add_run(v); set_font(r, size=9.5)
# 表头底色
for c in tbl.rows[0].cells:
    tcPr = c._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn('w:shd'), {qn('w:val'):'clear', qn('w:fill'):'DCE6F1'})
    tcPr.append(shd)
para("", size=4)

# ===== 二、推导依据 =====
heading("二、推导依据", 0)
heading("2.1 标准与规范依据", 2)
for t in BASIS["standard"]:
    para("· " + t, size=10)
heading("2.2 混合物危险组分浓度限值（关键判据）", 2)
for t in BASIS["concentration_limit"]:
    para("· " + t, size=10)
heading("2.3 产品身份依据（第1部分）", 2)
for t in BASIS["product"]:
    para("· " + t, size=10)
heading("2.4 成分依据（第3部分）", 2)
for t in BASIS["composition"]:
    para("· " + t, size=10)
heading("2.5 理化性质依据（第9部分）", 2)
for t in BASIS["physchem"]:
    para("· " + t, size=10)
heading("2.6 组分危害数据（网络检索，用于推导其余部分）", 2)
for t in BASIS["ingredient_hazards"]:
    para("· " + t, size=10)

# ===== 三、各Section推导 =====
heading("三、各部分的规范化推导（具体应写内容）", 0)
para("以下按 GB/T 16483-2008 的 16 个部分逐一给出：原文存在问题 → 推导后的规范应写内容。"
     "其中第1、3、9部分为已知事实（据原文档整理并规范化），第2、4~8、10~16部分均由前三部分推导得出。",
     size=10.5)

all_sections = SECTIONS_1_8 + SECTIONS_9_16
for s in all_sections:
    heading(f"{s['num']}  {s['title']}", 1)
    para("【原文问题】" + s["diagnosis"], size=10, color=(0x9C,0x27,0x00))
    para("【推导后的规范应写内容】", size=10.5, bold=True, cn='黑体')
    for line in s["content"]:
        para(line, size=10.5, space_after=2)

# ===== 四、参考资料 =====
heading("四、参考资料", 0)
for t in REFERENCES:
    para("· " + t, size=10)

# 保存
out = os.path.join(BASE, "OS-1330 MSDS规范化推导方案.docx")
doc.save(out)
print("saved:", out)
