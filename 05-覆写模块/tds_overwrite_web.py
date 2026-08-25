"""TDS 覆写工具的目录、统一字段和固定模板产出服务。

本模块只读扫描产品 TDS 资料目录，选中文件后按需解析，随后把字段写入正式版程序
自己的固定模板副本。原始 .doc/.docx/.pdf 和正式模板均不会被覆写。
"""
from __future__ import annotations

import hashlib
import re
import shutil
import tempfile
import threading
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path(r"F:\正式项目与模块化内容\冠志\MSDS")
WEB_ROOT = ROOT / "08-正式版程序"
SOURCE_ROOT = Path(r"F:\冠志工作空间\产品\TDS MSDS\TDS MSDS\产品 TDS MSDS -- WORD版本")
TEMPLATE_ROOT = WEB_ROOT / "tds_templates"
OUTPUT_ROOT = WEB_ROOT / "outputs"

_WPS_LOCK = threading.RLock()
_CATALOG_LOCK = threading.RLock()
_CATALOG: list[dict[str, Any]] | None = None
_CATALOG_BY_ID: dict[str, dict[str, Any]] = {}

COMPANY_PROFILES = {
    "guanzhi": {
        "label": "冠志",
        "name_zh": "广州冠志新材料科技有限公司",
        "name_en": "Guangzhou Guanzhi New Material Technology Co., Ltd.",
        "tel_zh": "电话：020-82567990\t传真：020-32214789",
        "tel_en": "Tel: +86 20 8256 7990\tFax: +86 20 3221 4789",
        "address_zh": "地址：广州市黄埔区科学城掬泉路 3 号广州国际企业孵化器 A 区 1106 室",
        "address_en": "Address: Room 1106, Block A, Guangzhou International Business Incubator, No. 3 Juquan Road, Science City, Huangpu District, Guangzhou, China",
    },
    "guocai": {
        "label": "国彩",
        "name_zh": "英德市国彩精细化工有限公司",
        "name_en": "Yingde Guocai Fine Chemical Co., Ltd.",
        "tel_zh": "电话：0763-2811205\t传真：0763-2811024",
        "tel_en": "Tel: +86 763 2811 205\tFax: +86 763 2811 024",
        "address_zh": "地址：广东省英德市白沙镇太平村更古坑凯迪工业园区",
        "address_en": "Address: Kaidi Industrial Park, Genggukeng, Taiping Village, Baisha Town, Yingde, Guangdong, China",
    },
}

HEADINGS = {
    "description": ("产品描述", "characterization", "product description", "description"),
    "performance": ("性能指标", "specification property", "technical data", "performance data", "specification"),
    "features": ("产品特性", "product features", "features", "property"),
    "application": ("应用", "application", "applications", "apllication"),
    "storage": ("储存", "storage", "storage condition", "storage conditions"),
    "supply": ("供货形式", "供货状态", "supply form", "delivery form", "physical form"),
    "packaging": ("包装", "包装规格", "packaging", "package"),
    "precaution": ("注意事项", "使用注意", "precaution", "handling", "notice"),
}

LABELS = {
    "外观": "Appearance", "颜色": "Color", "透明度": "Transparency", "固体份": "Solid content",
    "固含": "Solid content", "固含量": "Solid content", "粘度": "Viscosity", "PH值": "pH value",
    "pH值": "pH value", "密度": "Density", "羟基含量": "Hydroxyl content", "酸值": "Acid value",
    "胺值": "Amine value", "水分": "Moisture", "闪点": "Flash point", "细度": "Fineness",
    "玻璃化温度": "Glass transition temperature", "软化点": "Softening point", "熔点": "Melting point",
    "沸点": "Boiling point", "模量": "Modulus", "断裂伸长率": "Elongation at break",
    "拉伸强度": "Tensile strength", "溶剂": "Solvent", "活性成分": "Active content",
    "乳液外观": "Emulsion appearance", "环氧当量（EEW）": "Epoxy equivalent weight (EEW)",
    "环氧当量": "Epoxy equivalent weight", "PH值（25℃）": "pH value (25 °C)",
    "粘度（25℃）": "Viscosity (25 °C)",
    "项目": "Item", "指标": "Index", "单位": "Unit", "测试方法": "Test method",
    "透明至半透明液体": "Transparent to translucent liquid", "半透明液体": "Translucent liquid",
    "透明液体": "Transparent liquid", "乳白色液体": "Milky white liquid", "白色液体": "White liquid",
    "目测": "Visual inspection", "约": "Approx.", "建议储存在": "Store at", "通风干燥": "dry and well-ventilated",
    "严禁冰冻": "Protect from freezing", "保质期": "shelf life", "自发货之日起": "from the date of shipment",
}

TEXT_REPLACEMENTS = {
    "水性": "waterborne", "水分散型": "water-dispersible", "水分散性": "water-dispersible",
    "聚氨酯分散体": "polyurethane dispersion", "聚氨酯": "polyurethane", "丙烯酸": "acrylic",
    "聚酯": "polyester", "环氧树脂": "epoxy resin", "环氧": "epoxy", "环氧分散体": "epoxy dispersion",
    "水性金属防锈底漆": "waterborne metal anti-corrosion primers", "金属底材": "metal substrates",
    "防锈底漆": "anti-corrosion primers", "防护涂料": "protective coatings", "环氧当量": "epoxy equivalent weight",
    "自乳化固体环氧": "self-emulsifying solid epoxy", "储存稳定性佳": "excellent storage stability",
    "储存过程粘度稳定": "consistent viscosity during storage", "配方设计": "formulation design",
    "色浆研磨": "color paste grinding", "耐水性": "water resistance", "耐盐雾性": "salt spray resistance",
    "树脂": "resin", "固化剂": "curing agent", "涂料": "coating",
    "皮革": "leather", "纺织": "textile", "塑胶基材": "plastic substrates", "塑料基材": "plastic substrates",
    "高光泽": "high gloss", "丰满度": "fullness", "成膜": "film formation", "耐水性": "water resistance",
    "耐沸水": "boiling water resistance", "耐溶剂": "solvent resistance", "抗刮": "scratch resistance",
    "耐磨": "wear resistance", "抗冻融": "freeze-thaw resistance", "颜料润湿性": "pigment wetting",
    "剪切稳定性": "shear stability", "环保": "environmentally friendly", "不含": "free of",
    "无有机锡": "free of organotin", "固体份含量": "Solid content", "固含量": "Solid content",
    "粘度": "Viscosity", "密度": "Density", "外观": "Appearance", "溶剂": "Solvent",
    "储存": "Storage", "应用": "Application", "产品描述": "Characterization", "产品特性": "Product features",
}


def _clean(value: Any) -> str:
    return re.sub(r"[\r\n\x07\x0b]+", " ", str(value or "")).replace("  ", " ").strip()


def _normalize_heading(value: str) -> str:
    value = _clean(value).strip("【】[]:： ").lower()
    return re.sub(r"\s+", " ", value)


def _heading_kind(value: str) -> str | None:
    norm = _normalize_heading(value)
    for kind, labels in HEADINGS.items():
        if any(norm == x.lower() or norm.startswith(x.lower()) for x in labels):
            return kind
    return None


def _model_from_text(*values: str) -> str:
    joined = " ".join(values)
    matches = re.findall(r"(?<![A-Za-z0-9])([A-Za-z]{1,12}[-_]\d{2,7}[A-Za-z0-9+]*)(?![A-Za-z0-9])", joined)
    return matches[0].upper().replace("_", "-") if matches else ""


def _language_from_name(name: str) -> str:
    low = name.lower()
    if re.search(r"(?:tds|\b)[ _-]*(?:en|eng|english)|英文", low):
        return "en"
    if re.search(r"(?:tds|\b)[ _-]*(?:cn|zh|chinese)|中文", low):
        return "zh"
    return "zh" if re.search(r"[\u3400-\u9fff]", name) else "unknown"


def _company_from_name(name: str) -> str:
    low = name.lower()
    if "国彩" in name or "guocai" in low or "英德" in name:
        return "guocai"
    if "冠志" in name or "guanzhi" in low or "广州" in name:
        return "guanzhi"
    return "unknown"


def _model_key(item: dict[str, Any]) -> str:
    return (item.get("model") or Path(item["name"]).stem).upper().replace("_", "-")


def scan_catalog(refresh: bool = False) -> list[dict[str, Any]]:
    """建立轻量索引，不打开文档；所有源文件均以只读资料处理。"""
    global _CATALOG, _CATALOG_BY_ID
    with _CATALOG_LOCK:
        if _CATALOG is not None and not refresh:
            return deepcopy(_CATALOG)
        if not SOURCE_ROOT.is_dir():
            raise FileNotFoundError(f"TDS 源目录不存在：{SOURCE_ROOT}")
        rows: list[dict[str, Any]] = []
        for path in sorted(SOURCE_ROOT.rglob("*"), key=lambda x: str(x).lower()):
            if not path.is_file() or path.name.startswith("~$"):
                continue
            if path.suffix.lower() not in {".doc", ".docx", ".pdf"} or "tds" not in path.name.lower():
                continue
            rel = path.relative_to(SOURCE_ROOT)
            rel_text = str(rel)
            source_id = hashlib.sha1(rel_text.lower().encode("utf-8", "ignore")).hexdigest()[:20]
            category = rel.parts[0] if len(rel.parts) > 1 else "未分类"
            row = {
                "source_id": source_id, "name": path.name, "relative_path": rel_text,
                "path": str(path), "suffix": path.suffix.lower().lstrip("."),
                "language": _language_from_name(path.name), "company": _company_from_name(path.name),
                "model": _model_from_text(path.name, rel_text), "category": category,
                "editable": path.suffix.lower() in {".doc", ".docx"}, "size": path.stat().st_size,
                "modified": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            }
            rows.append(row)
        by_id = {row["source_id"]: row for row in rows}
        model_groups: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            model_groups.setdefault(_model_key(row), []).append(row)
        for row in rows:
            variants = model_groups.get(_model_key(row), [])
            row["variants"] = len(variants)
            row["counterparts"] = sum(1 for other in variants if other["source_id"] != row["source_id"] and other["language"] != row["language"])
        _CATALOG, _CATALOG_BY_ID = rows, by_id
        return deepcopy(rows)


def catalog(query: str = "", language: str = "", company: str = "", category: str = "", limit: int = 250) -> dict[str, Any]:
    rows = scan_catalog()
    q = str(query or "").strip().lower()
    filtered = [row for row in rows if (not q or q in " ".join(str(row.get(k, "")) for k in ("name", "model", "category", "relative_path")).lower())
                and (not language or row["language"] == language)
                and (not company or row["company"] == company)
                and (not category or row["category"] == category)]
    return {"items": filtered[:max(1, min(int(limit or 250), 1000))], "count": len(filtered),
            "total": len(rows), "categories": sorted({row["category"] for row in rows})}


def _resolve_source(source_id: str) -> dict[str, Any]:
    scan_catalog()
    row = _CATALOG_BY_ID.get(str(source_id or ""))
    if not row:
        raise FileNotFoundError("TDS 源文件不存在或不在允许的源目录内")
    path = Path(row["path"]).resolve()
    root = SOURCE_ROOT.resolve()
    if root not in path.parents or not path.is_file():
        raise ValueError("TDS 源文件路径越界或已不存在")
    return row


def _extract_word(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        paras = [_clean(p.text) for p in doc.paragraphs]
        tables = [[[ _clean(cell.text) for cell in row.cells] for row in table.rows] for table in doc.tables]
        return paras, tables
    with _WPS_LOCK:
        try:
            import pythoncom
            pythoncom.CoInitialize()
        except Exception:
            pythoncom = None
        app = None
        doc = None
        try:
            from win32com.client import DispatchEx
            app = DispatchEx("KWPS.Application")
            app.Visible = False
            doc = app.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
            paras = [_clean(doc.Paragraphs(i).Range.Text) for i in range(1, doc.Paragraphs.Count + 1)]
            tables: list[list[list[str]]] = []
            for ti in range(1, doc.Tables.Count + 1):
                table = doc.Tables(ti)
                rows = []
                for ri in range(1, table.Rows.Count + 1):
                    values = []
                    for ci in range(1, table.Columns.Count + 1):
                        values.append(_clean(table.Cell(ri, ci).Range.Text))
                    rows.append(values)
                tables.append(rows)
            return paras, tables
        finally:
            try:
                if doc is not None:
                    doc.Close(False)
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            if pythoncom is not None:
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass


def _extract_pdf(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    try:
        import fitz
        doc = fitz.open(path)
        text = "\n".join(page.get_text("text") for page in doc)
        return [_clean(line) for line in text.splitlines() if _clean(line)], []
    except Exception as exc:
        raise RuntimeError(f"PDF 读取失败：{path.name}：{exc}") from exc


def _read_raw(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    if path.suffix.lower() == ".pdf":
        return _extract_pdf(path)
    return _extract_word(path)


def _section_spans(paragraphs: list[str]) -> dict[str, tuple[int, int]]:
    marks = [(i, _heading_kind(text)) for i, text in enumerate(paragraphs) if _heading_kind(text)]
    spans: dict[str, tuple[int, int]] = {}
    for pos, (start, kind) in enumerate(marks):
        end = marks[pos + 1][0] if pos + 1 < len(marks) else len(paragraphs)
        spans.setdefault(kind, (start, end))
    return spans


def _paragraph_body(paragraphs: list[str], spans: dict[str, tuple[int, int]], kind: str) -> list[str]:
    if kind not in spans:
        return []
    start, end = spans[kind]
    return [text for text in paragraphs[start + 1:end] if text and not _heading_kind(text)]


def _table_performance(tables: list[list[list[str]]]) -> list[dict[str, str]]:
    best: list[list[str]] = []
    for table in tables:
        if not table:
            continue
        first = " ".join(table[0]).lower()
        if any(x in first for x in ("项目", "指标", "单位", "item", "index", "unit", "test method")):
            best = table
            break
        if len(table) > len(best):
            best = table
    if not best:
        return []
    header = [x.lower() for x in best[0]]
    header_joined = " ".join(header)
    has_header = any(x in header_joined for x in ("项目", "item", "指标", "index", "project", "metric"))
    rows = best[1:] if has_header else best
    result = []
    for values in rows:
        values = list(values) + [""] * (4 - len(values))
        item, index, unit, method = values[:4]
        if not any((item, index, unit, method)):
            continue
        result.append({"item": item, "index": index, "unit": unit, "method": method})
    return result


def _translate_text(text: str) -> tuple[str, list[str]]:
    source = _clean(text)
    if not source or not re.search(r"[\u3400-\u9fff]", source):
        return source, []
    result = source
    for term in sorted(LABELS, key=len, reverse=True):
        result = result.replace(term, LABELS[term])
    for term in sorted(TEXT_REPLACEMENTS, key=len, reverse=True):
        result = result.replace(term, TEXT_REPLACEMENTS[term])
    unresolved = sorted(set(re.findall(r"[\u3400-\u9fff]+", result)))
    return result, unresolved


def _translate_form(form: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    out = deepcopy(form)
    review: list[dict[str, Any]] = []
    for key in ("product_name", "description", "supply", "packaging", "precaution", "storage"):
        value, unresolved = _translate_text(out.get(key, ""))
        out[key] = value
        if unresolved:
            review.append({"field": key, "unresolved": unresolved, "source": form.get(key, "")})
    for key in ("features", "application"):
        values = []
        for index, item in enumerate(out.get(key) or []):
            value, unresolved = _translate_text(item)
            values.append(value)
            if unresolved:
                review.append({"field": f"{key}[{index}]", "unresolved": unresolved, "source": item})
        out[key] = values
    performance = []
    for index, row in enumerate(out.get("performance") or []):
        translated = dict(row)
        for key in ("item", "index", "unit", "method"):
            value, unresolved = _translate_text(translated.get(key, ""))
            translated[key] = value
            if unresolved:
                review.append({"field": f"performance[{index}].{key}", "unresolved": unresolved, "source": row.get(key, "")})
        performance.append(translated)
    out["performance"] = performance
    return out, review


def _parse_form(row: dict[str, Any]) -> dict[str, Any]:
    paragraphs, tables = _read_raw(Path(row["path"]))
    spans = _section_spans(paragraphs)
    model = row.get("model") or _model_from_text(*paragraphs[:12])
    title = ""
    for text in paragraphs[: max(1, spans.get("description", (len(paragraphs),))[0] if "description" in spans else 16)]:
        if text and not _heading_kind(text) and not re.search(r"(?:电话|传真|地址|tel|fax|address|co\.,? ltd)", text, re.I):
            if model and model.lower() in text.lower():
                title = text
                break
    if not title:
        title = next((x for x in paragraphs[6:16] if x and not _heading_kind(x)), model)
    form = {
        "source_id": row["source_id"], "source_name": row["name"], "source_language": row["language"],
        "source_company": row["company"], "category": row["category"], "model": model,
        "product_name": title, "description": (_paragraph_body(paragraphs, spans, "description") or [""])[0],
        "supply": (_paragraph_body(paragraphs, spans, "supply") or [""])[0],
        "packaging": (_paragraph_body(paragraphs, spans, "packaging") or [""])[0],
        "precaution": (_paragraph_body(paragraphs, spans, "precaution") or [""])[0],
        "features": _paragraph_body(paragraphs, spans, "features"),
        "application": _paragraph_body(paragraphs, spans, "application"),
        "storage": " ".join(_paragraph_body(paragraphs, spans, "storage")),
        "performance": _table_performance(tables),
    }
    if not form["performance"]:
        # Some PDFs/old Word files do not expose a real table. Keep the raw
        # performance body visible for manual review rather than silently dropping it.
        body = _paragraph_body(paragraphs, spans, "performance")
        if body:
            form["performance"] = [{"item": "Raw performance text", "index": "\n".join(body), "unit": "", "method": ""}]
    return form


def _variants_for(row: dict[str, Any]) -> list[dict[str, Any]]:
    key = _model_key(row)
    return [item for item in scan_catalog() if _model_key(item) == key]


def load_source(source_id: str, language: str = "zh", company: str = "guanzhi") -> dict[str, Any]:
    language = "en" if str(language).lower() == "en" else "zh"
    company = "guocai" if str(company).lower() == "guocai" else "guanzhi"
    selected = _resolve_source(source_id)
    form = _parse_form(selected)
    variants = _variants_for(selected)
    preferred = next((item for item in variants if item["language"] == language and item["company"] == company), None)
    if preferred is None:
        preferred = next((item for item in variants if item["language"] == language), None)
    review: list[dict[str, Any]] = []
    if preferred and preferred["source_id"] != selected["source_id"]:
        paired_form = _parse_form(preferred)
        for key in ("product_name", "description", "supply", "packaging", "precaution", "features", "application", "storage", "performance"):
            if paired_form.get(key):
                form[key] = paired_form[key]
        form["paired_source_id"] = preferred["source_id"]
        form["paired_source_name"] = preferred["name"]
    elif (language == "en" and selected["language"] != "en") or (language == "zh" and selected["language"] == "en"):
        if language == "en":
            form, review = _translate_form(form)
        else:
            review.append({"field": "source_language", "unresolved": ["未找到中文配对文件"], "source": selected["name"]})
    form["target_language"] = language
    form["target_company"] = company
    form["translation_review"] = review
    form["available_variants"] = [{k: item[k] for k in ("source_id", "name", "language", "company", "suffix", "relative_path")} for item in variants]
    return form


def _template_path(language: str, company: str) -> Path:
    language = "en" if language == "en" else "zh"
    company = "guocai" if company == "guocai" else "guanzhi"
    filename = f"template_{'EN' if language == 'en' else 'CN'}_{'国彩' if company == 'guocai' else '冠志'}.docx"
    path = TEMPLATE_ROOT / filename
    if not path.is_file():
        raise FileNotFoundError(f"固定 TDS 模板不存在：{path}")
    return path


def templates() -> dict[str, Any]:
    return {"root": str(TEMPLATE_ROOT), "items": [
        {"language": language, "company": company, "format": fmt, "path": str(_template_path(language, company)),
         "label": f"{'中文' if language == 'zh' else 'English'} · {COMPANY_PROFILES[company]['label']} · {fmt.upper()}"}
        for language in ("zh", "en") for company in ("guanzhi", "guocai") for fmt in ("docx", "pdf")
    ]}


def _set_text(paragraph, value: str) -> None:
    # Assigning paragraph.text retains the paragraph's style and direct paragraph
    # layout while replacing all runs, which is appropriate for fixed text slots.
    paragraph.text = str(value or "")


def _find_heading_index(doc, kind: str) -> int | None:
    for index, paragraph in enumerate(doc.paragraphs):
        if _heading_kind(paragraph.text) == kind:
            return index
    return None


def _set_section_body(doc, kind: str, values: list[str] | str, slots: int = 1) -> None:
    heading_index = _find_heading_index(doc, kind)
    if heading_index is None:
        return
    values = [values] if isinstance(values, str) else [str(x or "") for x in values]
    values = [x for x in values if x.strip()]
    next_heading = len(doc.paragraphs)
    for index in range(heading_index + 1, len(doc.paragraphs)):
        if _heading_kind(doc.paragraphs[index].text):
            next_heading = index
            break
    # The fixed templates deliberately leave blank paragraphs for vertical
    # spacing. Locate only the non-empty body slots, otherwise text would be
    # written into the spacer and the original sample wording would remain.
    targets = [index for index in range(heading_index + 1, next_heading)
               if doc.paragraphs[index].text.strip() and not _heading_kind(doc.paragraphs[index].text)]
    if not targets:
        return
    for offset, target in enumerate(targets[:slots]):
        _set_text(doc.paragraphs[target], values[offset] if offset < len(values) else "")
    if values and len(values) > slots:
        _set_text(doc.paragraphs[targets[min(slots, len(targets)) - 1]], "\n".join(values[slots - 1:]))


def _set_header(doc, language: str, company: str) -> None:
    profile = COMPANY_PROFILES[company]
    texts = [profile[f"name_{language}"], profile[f"tel_{language}"], profile[f"address_{language}"]]
    for index, text in zip((0, 2, 4), texts):
        if index < len(doc.paragraphs):
            _set_text(doc.paragraphs[index], text)


def _set_performance_table(doc, rows: list[dict[str, str]], language: str) -> None:
    if not doc.tables:
        return
    table = doc.tables[0]
    rows = [row for row in rows if any(str(row.get(key, "")).strip() for key in ("item", "index", "unit", "method"))]
    # EP-1704's approved skeleton is a four-column table with a real header
    # row. Keep the header as row 0 and only resize the data rows below it.
    header = (["项目", "指标", "单位", "测试方法"] if language == "zh"
              else ["Item", "Index", "Unit", "Test method"])
    while len(table.rows) < max(1, len(rows) + 1):
        table.add_row()
    while len(table.rows) > max(1, len(rows) + 1):
        table._tbl.remove(table.rows[-1]._tr)
    for column, value in enumerate(header):
        table.rows[0].cells[column].text = value
    for row_index, data in enumerate(rows):
        cells = table.rows[row_index + 1].cells
        values = [data.get("item"), data.get("index"), data.get("unit"), data.get("method")]
        for column, value in enumerate(values):
            cells[column].text = _clean(value)


def _normalize_form_for_output(form: dict[str, Any], language: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    result = deepcopy(form)
    review = list(result.get("translation_review") or [])
    if language == "en":
        translated, extra = _translate_form(result)
        # Paired English files are already preferred. Only replace fields that
        # still contain Chinese, preserving their professional sample wording.
        for key, value in translated.items():
            if key not in {"translation_review"} and value:
                if isinstance(value, str) and re.search(r"[\u3400-\u9fff]", value):
                    result[key] = value
                elif isinstance(value, list) and any(re.search(r"[\u3400-\u9fff]", str(x)) for x in value):
                    result[key] = value
        review.extend(extra)
    return result, review


def _safe_output_name(name: str, fmt: str) -> str:
    fmt = "pdf" if str(fmt).lower() == "pdf" else "docx"
    raw = Path(str(name or "").strip()).name
    raw = Path(raw).stem if raw else f"TDS_覆写_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    if not re.fullmatch(r"[\w\-. 一-龥（）()]+", raw):
        raise ValueError("TDS 输出名称含有不允许的字符")
    return f"{raw}.{fmt}"


def _write_docx(form: dict[str, Any], language: str, company: str, output: Path) -> None:
    doc = Document(_template_path(language, company))
    _set_header(doc, language, company)
    title = form.get("product_name") or form.get("model") or "TDS"
    # The fixed templates use a single title slot before the first section heading.
    desc_index = _find_heading_index(doc, "description")
    if desc_index is not None:
        title_index = next((i for i in range(desc_index - 1, -1, -1) if doc.paragraphs[i].text.strip()), None)
        if title_index is not None:
            _set_text(doc.paragraphs[title_index], title)
    heading_labels = {"description": "【产品描述】" if language == "zh" else "【Characterization】",
                      "performance": "【性能指标】" if language == "zh" else "【Specification Property】",
                      "features": "【产品特性】" if language == "zh" else "【Product Features】",
                      "application": "【应用】" if language == "zh" else "【Application】",
                      "storage": "【储存】" if language == "zh" else "【Storage】"}
    for kind, label in heading_labels.items():
        index = _find_heading_index(doc, kind)
        if index is not None:
            _set_text(doc.paragraphs[index], label)
    _set_section_body(doc, "description", form.get("description", ""), 1)
    _set_performance_table(doc, form.get("performance") or [], language)
    _set_section_body(doc, "features", form.get("features") or [], 2)
    application = list(form.get("application") or [])
    if form.get("supply"):
        application.append(("供货形式：" if language == "zh" else "Supply form: ") + str(form["supply"]))
    if form.get("packaging"):
        application.append(("包装：" if language == "zh" else "Packaging: ") + str(form["packaging"]))
    _set_section_body(doc, "application", application, 1)
    _set_section_body(doc, "storage", form.get("storage", ""), 1)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def _convert_pdf(docx_path: Path, pdf_path: Path) -> None:
    # Prefer the same WPS COM instance used for legacy .doc reading. Native
    # ExportAsFixedFormat keeps the EP-1704 table/layout stable and avoids the
    # WPS CLI process waiting indefinitely on an already-running WPS session.
    com_error = None
    with _WPS_LOCK:
        app = None
        doc = None
        try:
            import pythoncom
            pythoncom.CoInitialize()
        except Exception:
            pythoncom = None
        try:
            from win32com.client import DispatchEx
            app = DispatchEx("KWPS.Application")
            app.Visible = False
            doc = app.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
            doc.ExportAsFixedFormat(str(pdf_path), 17, False)
        except Exception as exc:
            com_error = exc
        finally:
            try:
                if doc is not None:
                    doc.Close(False)
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            if pythoncom is not None:
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
    if com_error is not None:
        from msds_multiformat import convert_docx_to_pdf
        convert_docx_to_pdf(docx_path, pdf_path)
    if not pdf_path.is_file() or pdf_path.stat().st_size < 100:
        detail = f"；COM错误：{com_error}" if com_error else ""
        raise RuntimeError(f"PDF 产出未生成有效文件：{pdf_path}{detail}")


def _apply_form_override(loaded: dict[str, Any], source_id: str, target_language: str,
                         company: str, form: dict[str, Any], form_language: str) -> dict[str, Any]:
    """Apply only actual user edits when a batch crosses language boundaries.

    The UI normally submits the Chinese form for all eight targets. Replacing
    the paired English sample wholesale would corrupt a clean English output.
    Compare with the original form in the submitted language, keep only edits,
    and translate those edits for the other language with review metadata.
    """
    allowed = ("model", "product_name", "description", "supply", "packaging", "precaution",
               "storage", "performance", "features", "application")
    form_language = "en" if str(form_language).lower() == "en" else "zh"
    if form_language == target_language:
        for key in allowed:
            if key in form:
                loaded[key] = deepcopy(form[key])
        return loaded
    reference = load_source(source_id, form_language, company)
    changed = {key: deepcopy(form[key]) for key in allowed
               if key in form and form.get(key) != reference.get(key)}
    if not changed:
        return loaded
    if target_language == "en" and form_language == "zh":
        translated, review = _translate_form(changed)
        for key, value in translated.items():
            if key in allowed:
                loaded[key] = value
        loaded.setdefault("translation_review", []).extend(review)
    else:
        for key, value in changed.items():
            loaded[key] = value
        loaded.setdefault("translation_review", []).append({
            "field": "form_language", "unresolved": ["跨语言表单编辑需要人工翻译"], "source": form_language,
        })
    return loaded


def generate(source_id: str, language: str, company: str, output_format: str = "docx",
             form: dict[str, Any] | None = None, output_name: str = "",
             form_language: str = "zh") -> dict[str, Any]:
    language = "en" if str(language).lower() == "en" else "zh"
    company = "guocai" if str(company).lower() == "guocai" else "guanzhi"
    output_format = "pdf" if str(output_format).lower() == "pdf" else "docx"
    loaded = load_source(source_id, language, company)
    if form:
        loaded = _apply_form_override(loaded, source_id, language, company, form, form_language)
    final_form, review = _normalize_form_for_output(loaded, language)
    if len(final_form.get("features") or []) > 2:
        review.append({"field": "features", "unresolved": ["固定 EP-1704 骨架只有 2 个特性槽位"],
                       "source": final_form.get("features")})
    if len(final_form.get("application") or []) > 1:
        review.append({"field": "application", "unresolved": ["固定 EP-1704 骨架只有 1 个应用槽位"],
                       "source": final_form.get("application")})
    model = final_form.get("model") or _model_from_text(final_form.get("product_name", "")) or "TDS"
    default_stem = f"{model}_{'冠志' if company == 'guanzhi' else '国彩'}_{'CN' if language == 'zh' else 'EN'}"
    filename = _safe_output_name(output_name or default_stem, output_format)
    output = OUTPUT_ROOT / filename
    if output_format == "docx":
        _write_docx(final_form, language, company, output)
    else:
        with tempfile.NamedTemporaryFile(prefix="tds_render_", suffix=".docx", dir=OUTPUT_ROOT, delete=False) as handle:
            temp_docx = Path(handle.name)
        try:
            _write_docx(final_form, language, company, temp_docx)
            _convert_pdf(temp_docx, output)
        finally:
            try:
                temp_docx.unlink(missing_ok=True)
            except OSError:
                pass
    return {"ok": True, "output": str(output), "output_name": output.name,
            "download_url": f"/api/tds/download?name={output.name}", "source_id": source_id,
            "source_name": loaded.get("source_name", ""), "model": model, "language": language,
            "company": company, "output_format": output_format, "translation_review": review,
            "template": str(_template_path(language, company))}


def batch(source_id: str, form: dict[str, Any] | None = None, output_name: str = "",
          form_language: str = "zh") -> dict[str, Any]:
    loaded = load_source(source_id, "zh", "guanzhi")
    model = loaded.get("model") or "TDS"
    stem = Path(str(output_name or model)).stem
    if not re.fullmatch(r"[\w\-. 一-龥（）()]+", stem):
        raise ValueError("TDS 批量输出名称含有不允许的字符")
    results = []
    for language in ("zh", "en"):
        for company in ("guanzhi", "guocai"):
            for output_format in ("docx", "pdf"):
                label = f"{stem}_{'冠志' if company == 'guanzhi' else '国彩'}_{'中文' if language == 'zh' else '英文'}_{'Word' if output_format == 'docx' else 'PDF'}"
                try:
                    results.append(generate(source_id, language, company, output_format, form, label, form_language))
                except Exception as exc:
                    results.append({"ok": False, "source_id": source_id, "language": language,
                                    "company": company, "output_format": output_format,
                                    "output_name": f"{label}.{output_format}", "error": str(exc)})
    ok_items = [item for item in results if item.get("ok") and Path(item.get("output", "")).is_file()]
    archive = None
    if len(ok_items) == 8:
        archive = OUTPUT_ROOT / _safe_output_name(f"{stem}_TDS八件套", "zip")
        archive = archive.with_suffix(".zip")
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
            for item in ok_items:
                bundle.write(item["output"], arcname=Path(item["output"]).name)
    return {"ok": len(ok_items) == 8, "outputs": results, "archive_name": archive.name if archive else "",
            "download_url": f"/api/tds/batch-download?name={archive.name}" if archive else "",
            "error": "" if len(ok_items) == 8 else f"8 个目标中成功 {len(ok_items)} 个"}


def preview(source_id: str, language: str, company: str, form: dict[str, Any] | None = None) -> dict[str, Any]:
    loaded = load_source(source_id, language, company)
    if form:
        for key, value in form.items():
            if key in loaded:
                loaded[key] = value
    final_form, review = _normalize_form_for_output(loaded, language)
    if len(final_form.get("features") or []) > 2:
        review.append({"field": "features", "unresolved": ["固定 EP-1704 骨架只有 2 个特性槽位"],
                       "source": final_form.get("features")})
    if len(final_form.get("application") or []) > 1:
        review.append({"field": "application", "unresolved": ["固定 EP-1704 骨架只有 1 个应用槽位"],
                       "source": final_form.get("application")})
    return {"ok": True, "form": final_form, "translation_review": review,
            "template": str(_template_path(language, company)),
            "company_profile": COMPANY_PROFILES[company]}


def safe_output(name: str) -> Path:
    path = (OUTPUT_ROOT / Path(str(name or "")).name).resolve()
    root = OUTPUT_ROOT.resolve()
    if root not in path.parents or not path.is_file():
        raise FileNotFoundError("TDS 产出文件不存在")
    return path
