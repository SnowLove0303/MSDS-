"""CAS Section 2 result library initializer, NMP seed importer, and validator.

This module deliberately keeps CAS identity data, legal source rules, and
model-level inference results in separate stores.  It imports the NMP workbook
as a structured example without treating blank summary cells as no-hazard.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - runtime dependency guard
    raise SystemExit("缺少 openpyxl，无法读取 NMP Excel 模板。") from exc


SCRIPT_DIR = Path(__file__).resolve().parent
SCHEMA_PATH = SCRIPT_DIR / "cas_section2_result_schema.sql"
DEFAULT_DB = Path(r"F:\正式项目与模块化内容\冠志\MSDS\03-数据库\正式库\Data Base\cas_section2_result.db")
DEFAULT_TEMPLATE = Path(r"F:\正式项目与模块化内容\冠志\MSDS\04-推断引擎\数据库模板\N-甲基-2-吡咯烷酮-NMP-安全数据表.xlsx")
CAS_EXPECTED = "872-50-4"

STANDARD_SKELETON_FIELDS = (
    ("standard_name", "标准名称"),
    ("cas_no", "CAS"),
    ("ec_no", "EC"),
    ("index_no", "Index No"),
    ("ghs_classification", "GHS 危险性类别"),
    ("pictogram", "象形图"),
    ("signal_word", "信号词"),
    ("hazard_statement", "危险性说明"),
    ("precautionary_statement", "防范说明"),
    ("physical_chemical_hazard", "物理与化学危险"),
    ("health_hazard", "健康危害"),
    ("environmental_hazard", "环境危害"),
    ("other_hazard", "其他危害"),
)

# Section 2 standard results are final output values, not evidence notes.
# Evidence/source wording is kept in the evidence tables and must not leak
# into the 13-field result contract consumed by Web and MSDS generation.
STANDARD_RESULT_FORBIDDEN_MARKERS = (
    "输出应保留来源和方法",
    "来源和方法",
    "追溯分类证据",
    "国际参考事实",
    "需结合具体配方和暴露场景人工复核",
    "来源范围",
    "参考事实",
    "参考数据",
    "本次不将",
    "未据此自动",
    "暂无可直接升级",
    "仅作为混合物引用事实",
    "可追溯物质级分类",
    "人工复核",
)


def validate_standard_result_text(field_key: str, value_text: str) -> str:
    """Allow only direct Section 2 output text in the fixed result layer."""
    text = str(value_text or "").strip() or "无数据"
    if text == "无数据":
        return text
    matched = [marker for marker in STANDARD_RESULT_FORBIDDEN_MARKERS if marker in text]
    if matched:
        raise ValueError(
            f"标准结果字段 {field_key} 含有证据/说明性措辞，不允许写入结果骨架：{', '.join(matched)}"
        )
    return text

H_TEXT_RE = re.compile(r"^(H\d{3}[A-Z]?(?:\+H\d{3}[A-Z]?)*)\s+(.+)$")
P_TEXT_RE = re.compile(r"^(P\d{3}[A-Z]?(?:\+P\d{3}[A-Z]?)*)\s+(.+)$")
H_CODE_RE = re.compile(r"(?<![A-Za-z])H\d{3}[A-Z]?(?:\+H\d{3}[A-Z]?)*(?![A-Za-z0-9])")

DOMAIN_MAP = {
    "生殖毒性": "reproductive_toxicity",
    "特定靶器官毒性——单次接触": "specific_target_organ_toxicity_single_exposure",
    "皮肤刺激": "skin_irritation",
    "眼睛刺激": "eye_irritation",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(db_path))
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def ensure_migrated_columns(connection: sqlite3.Connection) -> None:
    """Add columns introduced after the initial NMP template import."""
    existing = {
        row[1]
        for row in connection.execute("PRAGMA table_info(cas_s2_source)")
    }
    migrations = {
        "source_uri": "ALTER TABLE cas_s2_source ADD COLUMN source_uri TEXT NOT NULL DEFAULT ''",
        "clause_reference": "ALTER TABLE cas_s2_source ADD COLUMN clause_reference TEXT NOT NULL DEFAULT ''",
        "source_priority": "ALTER TABLE cas_s2_source ADD COLUMN source_priority TEXT NOT NULL DEFAULT 'reference_only'",
    }
    for column, statement in migrations.items():
        if existing and column not in existing:
            connection.execute(statement)


def initialize(db_path: Path) -> None:
    if not SCHEMA_PATH.is_file():
        raise FileNotFoundError(f"缺少数据库结构文件：{SCHEMA_PATH}")
    with connect(db_path) as connection:
        ensure_migrated_columns(connection)
        connection.execute("DROP VIEW IF EXISTS v_cas_s2_current_profile")
        connection.execute("DROP VIEW IF EXISTS v_cas_s2_current_standard_result")
        connection.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))


def reset_database(db_path: Path) -> None:
    """Reset only this library after an external backup has been made."""
    if not db_path.is_file():
        initialize(db_path)
        return
    with connect(db_path) as connection:
        connection.executescript(
            """
            DROP VIEW IF EXISTS v_cas_s2_current_profile;
            DROP VIEW IF EXISTS v_cas_s2_current_standard_result;
            DROP TABLE IF EXISTS cas_s2_result_field;
            DROP TABLE IF EXISTS cas_s2_skeleton_field;
            DROP TABLE IF EXISTS cas_s2_review;
            DROP TABLE IF EXISTS cas_s2_hazard_summary;
            DROP TABLE IF EXISTS cas_s2_label_element;
            DROP TABLE IF EXISTS cas_s2_classification;
            DROP TABLE IF EXISTS cas_s2_condition;
            DROP TABLE IF EXISTS cas_s2_fact;
            DROP TABLE IF EXISTS cas_s2_source;
            DROP TABLE IF EXISTS cas_s2_profile;
            DROP TABLE IF EXISTS schema_meta;
            """
        )
        connection.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_locator(source_key: str, source_file_path: str, source_uri: str) -> tuple[str, str]:
    """Return a stable locator and hash for the legacy source uniqueness key."""
    local_path = Path(source_file_path) if source_file_path else None
    if local_path and local_path.is_file():
        return str(local_path), sha256_file(local_path)
    # The legacy table requires a non-empty uniqueness locator for remote rows.
    # The actual URL remains in source_uri; this marker is not counted as a file.
    return f"remote://{source_key}", ""


def upsert_source(connection: sqlite3.Connection, profile_id: int, source: dict[str, Any]) -> int:
    locator, source_hash = source_locator(
        source["key"], source.get("source_file_path", ""), source.get("source_uri", "")
    )
    existing = connection.execute(
        "SELECT source_id FROM cas_s2_source WHERE profile_id = ? AND source_file_path = ? AND source_hash_sha256 = ?",
        (profile_id, locator, source_hash),
    ).fetchone()
    values = (
        source.get("source_type", "reference"),
        source.get("source_title", source["key"]),
        source.get("source_file_name", Path(locator).name if not locator.startswith("remote://") else source["key"]),
        source.get("source_version", ""),
        source.get("effective_date", ""),
        source.get("collected_at", utc_now()),
        source.get("jurisdiction", "CN"),
        source.get("source_uri", ""),
        source.get("clause_reference", ""),
        source.get("source_priority", "reference_only"),
        source.get("source_status", "reference_only"),
        source.get("evidence_text", ""),
    )
    if existing:
        source_id = int(existing["source_id"])
        connection.execute(
            """
            UPDATE cas_s2_source
            SET source_type = ?, source_title = ?, source_file_name = ?, source_version = ?,
                effective_date = ?, collected_at = ?, jurisdiction = ?, source_uri = ?,
                clause_reference = ?, source_priority = ?, source_status = ?, evidence_text = ?
            WHERE source_id = ?
            """,
            (*values, source_id),
        )
        return source_id
    cursor = connection.execute(
        """
        INSERT INTO cas_s2_source (
            profile_id, source_type, source_title, source_file_path, source_file_name,
            source_hash_sha256, source_version, effective_date, collected_at, jurisdiction,
            source_uri, clause_reference, source_priority, source_status, evidence_text
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (profile_id, values[0], values[1], locator, values[2], source_hash, *values[3:]),
    )
    return int(cursor.lastrowid)


def upsert_condition(connection: sqlite3.Connection, profile_id: int, condition: dict[str, Any]) -> int:
    existing = connection.execute(
        "SELECT condition_id FROM cas_s2_condition WHERE profile_id = ? AND condition_key = ?",
        (profile_id, condition["condition_key"]),
    ).fetchone()
    values = (
        condition.get("applies_to_type", "profile"),
        condition.get("applies_to_code", "*"),
        condition["condition_type"],
        condition.get("lower_operator", ""),
        condition.get("lower_value"),
        condition.get("upper_operator", ""),
        condition.get("upper_value"),
        condition.get("unit", "%"),
        condition.get("scope", ""),
        condition.get("raw_text", ""),
        condition.get("status", "manual_review"),
        condition.get("priority", 100),
        condition.get("source_id"),
        condition.get("notes", ""),
    )
    if existing:
        condition_id = int(existing["condition_id"])
        connection.execute(
            """
            UPDATE cas_s2_condition
            SET applies_to_type = ?, applies_to_code = ?, condition_type = ?, lower_operator = ?,
                lower_value = ?, upper_operator = ?, upper_value = ?, unit = ?, scope = ?,
                raw_text = ?, status = ?, priority = ?, source_id = ?, notes = ?
            WHERE condition_id = ?
            """,
            (*values, condition_id),
        )
        return condition_id
    cursor = connection.execute(
        """
        INSERT INTO cas_s2_condition (
            profile_id, condition_key, applies_to_type, applies_to_code, condition_type,
            lower_operator, lower_value, upper_operator, upper_value, unit, scope, raw_text,
            status, priority, source_id, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (profile_id, condition["condition_key"], *values),
    )
    return int(cursor.lastrowid)


def upsert_fact(connection: sqlite3.Connection, profile_id: int, fact: dict[str, Any]) -> None:
    columns = (
        "fact_domain", "fact_code", "fact_label_zh", "value_text", "numeric_value",
        "lower_value", "upper_value", "unit", "qualifier", "applicability_scope",
        "value_status", "source_id", "clause_reference", "raw_text", "display_order",
    )
    values = tuple(fact.get(column, "" if column not in {"numeric_value", "lower_value", "upper_value"} else None) for column in columns)
    existing = connection.execute(
        "SELECT fact_id FROM cas_s2_fact WHERE profile_id = ? AND fact_code = ? AND source_id IS ? AND qualifier = ?",
        (profile_id, fact["fact_code"], fact.get("source_id"), fact.get("qualifier", "")),
    ).fetchone()
    if existing:
        connection.execute(
            "UPDATE cas_s2_fact SET " + ", ".join(f"{column} = ?" for column in columns) + " WHERE fact_id = ?",
            (*values, int(existing["fact_id"])),
        )
        return
    connection.execute(
        "INSERT INTO cas_s2_fact (profile_id, " + ", ".join(columns) + ") VALUES (?, " + ", ".join("?" for _ in columns) + ")",
        (profile_id, *values),
    )


def ensure_standard_result_fields(connection: sqlite3.Connection, profile_id: int) -> None:
    """Create all 13 fixed output slots; missing values are explicitly 无数据."""
    connection.execute(
        """
        INSERT OR IGNORE INTO cas_s2_result_field (
            profile_id, field_key, value_text, value_status, display_order
        )
        SELECT ?, field_key, '无数据', 'unknown', display_order
        FROM cas_s2_skeleton_field
        ORDER BY display_order
        """,
        (profile_id,),
    )


def upsert_standard_result_field(
    connection: sqlite3.Connection,
    profile_id: int,
    field_key: str,
    value_text: str,
    *,
    value_status: str = "unknown",
    condition_id: int | None = None,
    source_id: int | None = None,
    raw_text: str = "",
    display_order: int | None = None,
) -> None:
    """Write one canonical field without allowing free-form output fields."""
    field = connection.execute(
        "SELECT display_order FROM cas_s2_skeleton_field WHERE field_key = ?", (field_key,)
    ).fetchone()
    if field is None:
        raise ValueError(f"不允许写入标准骨架之外的字段：{field_key}")
    order = int(display_order if display_order is not None else field[0])
    text = validate_standard_result_text(field_key, value_text)
    connection.execute(
        """
        INSERT INTO cas_s2_result_field (
            profile_id, field_key, value_text, value_status, condition_id,
            source_id, raw_text, display_order
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(profile_id, field_key) DO UPDATE SET
            value_text = excluded.value_text,
            value_status = excluded.value_status,
            condition_id = excluded.condition_id,
            source_id = excluded.source_id,
            raw_text = excluded.raw_text,
            display_order = excluded.display_order
        """,
        (profile_id, field_key, text, value_status, condition_id, source_id, raw_text or text, order),
    )


def refresh_standard_result_fields(
    connection: sqlite3.Connection,
    profile_id: int,
    standard_summaries: dict[str, str] | None = None,
) -> None:
    """Build the exact 13-field output; evidence summaries never leak into it."""
    ensure_standard_result_fields(connection, profile_id)
    profile = connection.execute(
        "SELECT standard_name_zh, cas_no, ec_no, index_no FROM cas_s2_profile WHERE profile_id = ?",
        (profile_id,),
    ).fetchone()
    if profile is None:
        raise ValueError(f"找不到结果档案：{profile_id}")

    def write_identity(field_key: str, value: str) -> None:
        text = str(value or "").strip() or "无数据"
        upsert_standard_result_field(
            connection, profile_id, field_key, text,
            value_status="confirmed" if text != "无数据" else "unknown",
            raw_text=text,
        )

    write_identity("standard_name", profile["standard_name_zh"])
    write_identity("cas_no", profile["cas_no"])
    write_identity("ec_no", profile["ec_no"])
    write_identity("index_no", profile["index_no"])

    intrinsic = connection.execute(
        "SELECT condition_id FROM cas_s2_condition WHERE profile_id = ? AND condition_key = 'substance_intrinsic'",
        (profile_id,),
    ).fetchone()
    intrinsic_id = int(intrinsic[0]) if intrinsic else None

    def status_for(rows: list[sqlite3.Row], default: str = "unknown") -> str:
        if not rows:
            return default
        statuses = {str(row["value_status"] or "unknown") for row in rows}
        if "conflicting" in statuses:
            return "conflicting"
        if "derived_candidate" in statuses:
            return "derived_candidate"
        if "reference_only" in statuses and statuses <= {"reference_only"}:
            return "reference_only"
        if "confirmed" in statuses and statuses <= {"confirmed", "reference_only"}:
            return "confirmed"
        return "unknown"

    def first_source(rows: list[sqlite3.Row]) -> int | None:
        for row in rows:
            if row["source_id"] is not None:
                return int(row["source_id"])
        return None

    classifications = connection.execute(
        "SELECT classification_text, h_code, value_status, source_id FROM cas_s2_classification WHERE profile_id = ? ORDER BY display_order, classification_id",
        (profile_id,),
    ).fetchall()
    classification_text = "\n".join(
        f"{row['classification_text']}；{row['h_code']}" if row["h_code"] else str(row["classification_text"])
        for row in classifications
    ) or "无数据"
    upsert_standard_result_field(
        connection, profile_id, "ghs_classification", classification_text,
        value_status=status_for(classifications), condition_id=intrinsic_id,
        source_id=first_source(classifications), raw_text=classification_text,
    )

    labels = connection.execute(
        "SELECT element_type, element_code, element_text_zh, value_status, source_id FROM cas_s2_label_element WHERE profile_id = ? ORDER BY display_order, element_id",
        (profile_id,),
    ).fetchall()
    pictograms = [row for row in labels if row["element_type"] == "pictogram" and row["element_code"]]
    signal_words = [row for row in labels if row["element_type"] == "signal_word" and row["element_text_zh"]]
    hazard_statements = [row for row in labels if row["element_type"] == "hazard_statement"]
    precautionary_statements = [row for row in labels if row["element_type"] == "precautionary_statement"]
    upsert_standard_result_field(
        connection, profile_id, "pictogram", "；".join(row["element_code"] for row in pictograms) or "无数据",
        value_status=status_for(pictograms), condition_id=intrinsic_id,
        source_id=first_source(pictograms), raw_text="；".join(row["element_code"] for row in pictograms) or "无数据",
    )
    upsert_standard_result_field(
        connection, profile_id, "signal_word", "；".join(row["element_text_zh"] for row in signal_words) or "无数据",
        value_status=status_for(signal_words), condition_id=intrinsic_id,
        source_id=first_source(signal_words), raw_text="；".join(row["element_text_zh"] for row in signal_words) or "无数据",
    )

    def label_text(rows: list[sqlite3.Row]) -> str:
        return "\n".join(
            f"{row['element_code']} {row['element_text_zh']}".strip()
            for row in rows
            if row["element_code"] or row["element_text_zh"]
        ) or "无数据"

    for field_key, rows in (("hazard_statement", hazard_statements), ("precautionary_statement", precautionary_statements)):
        text = label_text(rows)
        upsert_standard_result_field(
            connection, profile_id, field_key, text, value_status=status_for(rows),
            condition_id=intrinsic_id, source_id=first_source(rows), raw_text=text,
        )

    standard_summaries = standard_summaries or {}
    for field_key in (
        "physical_chemical_hazard",
        "health_hazard",
        "environmental_hazard",
        "other_hazard",
    ):
        text = validate_standard_result_text(field_key, standard_summaries.get(field_key, "无数据"))
        upsert_standard_result_field(
            connection, profile_id, field_key, text,
            value_status="confirmed" if text != "无数据" else "unknown",
            condition_id=intrinsic_id if text != "无数据" else None,
            source_id=None,
            raw_text=text,
        )


def read_template(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"找不到 NMP 模板：{path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    if "Sheet1" not in workbook.sheetnames:
        raise ValueError("NMP 模板缺少 Sheet1。")

    values: dict[str, tuple[str, str]] = {}
    for row in workbook["Sheet1"].iter_rows(values_only=True):
        label = str(row[0] or "").strip().rstrip(":：")
        if not label:
            continue
        value = "" if row[1] is None else str(row[1]).strip()
        condition = "" if len(row) < 3 or row[2] is None else str(row[2]).strip()
        values[label] = (value, condition)

    required = {
        "标准名称", "CAS", "EC", "Index No", "GHS 危险性类别", "象形图", "信号词",
        "危险性说明", "防范说明", "物理与化学危险", "健康危害", "环境危害", "其他危害",
    }
    missing = sorted(required - values.keys())
    if missing:
        raise ValueError(f"NMP 模板缺少字段：{', '.join(missing)}")
    if values["CAS"][0] != CAS_EXPECTED:
        raise ValueError(f"模板 CAS 不是 {CAS_EXPECTED}，实际为：{values['CAS'][0]}")

    return {
        "identity": {
            "standard_name_zh": values["标准名称"][0],
            "cas_no": values["CAS"][0],
            "ec_no": values["EC"][0],
            "index_no": values["Index No"][0],
        },
        "classification_text": values["GHS 危险性类别"][0],
        "pictogram_text": values["象形图"][0],
        "signal_word": values["信号词"][0],
        "hazard_statement_text": values["危险性说明"][0],
        "precautionary_statement_text": values["防范说明"][0],
        "summary_values": {
            "physical_chemical": values.get("物理与化学危险", ("", ""))[0],
            "health": values.get("健康危害", ("", ""))[0],
            "environment": values.get("环境危害", ("", ""))[0],
            "other": values.get("其他危害", ("", ""))[0],
        },
        "condition_texts": sorted({condition for _, condition in values.values() if condition}),
        "source_rows": {label: condition for label, (_, condition) in values.items()},
    }


def classification_parts(line: str) -> tuple[str, str, str, str]:
    text = line.strip()
    match = H_CODE_RE.search(text)
    if not match:
        return "other", text, "", ""
    h_code = match.group(0)
    classification_text = text[: match.start()].rstrip("；; ")
    domain = "other"
    for prefix, mapped in DOMAIN_MAP.items():
        if classification_text.startswith(prefix):
            domain = mapped
            break
    category_match = re.search(r"类别\s*([0-9A-Za-z]+)", classification_text)
    category = category_match.group(1) if category_match else ""
    return domain, classification_text, category, h_code


def split_codes(text: str) -> list[str]:
    return [part for part in re.split(r"[；;,\s]+", text.strip()) if part]


def parse_code_lines(text: str, pattern: re.Pattern[str]) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        match = pattern.match(line)
        if match:
            rows.append((match.group(1), match.group(2), line))
        else:
            rows.append(("", line, line))
    return rows


def insert_nmp(db_path: Path, template_path: Path) -> dict[str, Any]:
    data = read_template(template_path)
    source_hash = sha256_file(template_path)
    now = utc_now()
    identity = data["identity"]

    initialize(db_path)
    with connect(db_path) as connection:
        existing = connection.execute(
            "SELECT profile_id FROM cas_s2_profile WHERE cas_no = ? AND is_current = 1",
            (identity["cas_no"],),
        ).fetchone()
        if existing:
            raise ValueError(f"当前库已存在 CAS {identity['cas_no']} 的当前档案，拒绝覆盖。")

        profile_cursor = connection.execute(
            """
            INSERT INTO cas_s2_profile (
                cas_no, standard_name_zh, ec_no, index_no, profile_version,
                profile_status, is_current, remarks, created_at, updated_at
            ) VALUES (?, ?, ?, ?, '1.0.0', 'draft', 1, ?, ?, ?)
            """,
            (
                identity["cas_no"],
                identity["standard_name_zh"],
                identity["ec_no"],
                identity["index_no"],
                "NMP 示例档案；由用户提供的 Excel 结构模板导入，法规核验状态待补充。",
                now,
                now,
            ),
        )
        profile_id = int(profile_cursor.lastrowid)

        source_cursor = connection.execute(
            """
            INSERT INTO cas_s2_source (
                profile_id, source_type, source_title, source_file_path,
                source_file_name, source_hash_sha256, source_status,
                collected_at, evidence_text
            ) VALUES (?, 'template_workbook', ?, ?, ?, ?, 'imported', ?, ?)
            """,
            (
                profile_id,
                template_path.stem,
                str(template_path),
                template_path.name,
                source_hash,
                now,
                "Sheet1 作为 CAS Section 2 结果骨架示例；空白危害摘要不解释为无危害。",
            ),
        )
        source_id = int(source_cursor.lastrowid)

        condition_text = data["condition_texts"][0] if data["condition_texts"] else ""
        condition_cursor = connection.execute(
            """
            INSERT INTO cas_s2_condition (
                profile_id, condition_key, applies_to_type, applies_to_code,
                condition_type, scope, raw_text, status, source_id, notes
            ) VALUES (?, 'profile_default', 'profile', '*', 'not_provided', 'mixture_component', ?, 'unknown', ?, ?)
            """,
            (
                profile_id,
                condition_text,
                source_id,
                "模板保留了含量触及条件字段，但当前没有数值、单位比较符或法规阈值；不得直接参与自动分类。",
            ),
        )
        condition_id = int(condition_cursor.lastrowid)

        for order, line in enumerate(data["classification_text"].splitlines(), start=1):
            if not line.strip():
                continue
            domain, classification_text, category, h_code = classification_parts(line)
            connection.execute(
                """
                INSERT INTO cas_s2_classification (
                    profile_id, hazard_domain, hazard_class, category,
                    classification_text, h_code, condition_id, value_status,
                    display_order, source_id, raw_text
                ) VALUES (?, ?, ?, ?, ?, ?, ?, 'derived_candidate', ?, ?, ?)
                """,
                (
                    profile_id,
                    domain,
                    classification_text.split("，")[0] if "，" in classification_text else classification_text,
                    category,
                    classification_text,
                    h_code,
                    condition_id,
                    order,
                    source_id,
                    line,
                ),
            )

        display_order = 1
        for code in split_codes(data["pictogram_text"]):
            connection.execute(
                """
                INSERT INTO cas_s2_label_element (
                    profile_id, element_type, element_code, element_text_zh,
                    condition_id, value_status, display_order, source_id, raw_text
                ) VALUES (?, 'pictogram', ?, '', ?, 'derived_candidate', ?, ?, ?)
                """,
                (profile_id, code, condition_id, display_order, source_id, code),
            )
            display_order += 1

        if data["signal_word"]:
            connection.execute(
                """
                INSERT INTO cas_s2_label_element (
                    profile_id, element_type, element_code, element_text_zh,
                    condition_id, value_status, display_order, source_id, raw_text
                ) VALUES (?, 'signal_word', '', ?, ?, 'derived_candidate', ?, ?, ?)
                """,
                (profile_id, data["signal_word"], condition_id, display_order, source_id, data["signal_word"]),
            )
            display_order += 1

        for code, text, raw in parse_code_lines(data["hazard_statement_text"], H_TEXT_RE):
            connection.execute(
                """
                INSERT INTO cas_s2_label_element (
                    profile_id, element_type, element_code, element_text_zh,
                    condition_id, value_status, display_order, source_id, raw_text
                ) VALUES (?, 'hazard_statement', ?, ?, ?, 'derived_candidate', ?, ?, ?)
                """,
                (profile_id, code, text, condition_id, display_order, source_id, raw),
            )
            display_order += 1

        for code, text, raw in parse_code_lines(data["precautionary_statement_text"], P_TEXT_RE):
            connection.execute(
                """
                INSERT INTO cas_s2_label_element (
                    profile_id, element_type, element_code, element_text_zh,
                    condition_id, value_status, display_order, source_id, raw_text
                ) VALUES (?, 'precautionary_statement', ?, ?, ?, 'derived_candidate', ?, ?, ?)
                """,
                (profile_id, code, text, condition_id, display_order, source_id, raw),
            )
            display_order += 1

        for domain, summary in data["summary_values"].items():
            connection.execute(
                """
                INSERT INTO cas_s2_hazard_summary (
                    profile_id, hazard_domain, summary_text, value_status,
                    condition_id, source_id, raw_text
                ) VALUES (?, ?, ?, 'unknown', ?, ?, ?)
                """,
                (profile_id, domain, summary, condition_id, source_id, summary),
            )

        connection.execute(
            """
            INSERT INTO cas_s2_review (
                profile_id, review_scope, review_status, review_note
            ) VALUES (?, 'profile', 'manual_review', ?)
            """,
            (
                profile_id,
                "NMP 示例已按模板导入；需要后续补充法规来源、分类依据和含量触发条件后，才能进入正式确认状态。",
            ),
        )
        refresh_standard_result_fields(connection, profile_id)

        return {
            "profile_id": profile_id,
            "cas_no": identity["cas_no"],
            "source_hash_sha256": source_hash,
            "classification_count": connection.execute(
                "SELECT COUNT(*) FROM cas_s2_classification WHERE profile_id = ?", (profile_id,)
            ).fetchone()[0],
            "label_element_count": connection.execute(
                "SELECT COUNT(*) FROM cas_s2_label_element WHERE profile_id = ?", (profile_id,)
            ).fetchone()[0],
            "summary_count": connection.execute(
                "SELECT COUNT(*) FROM cas_s2_hazard_summary WHERE profile_id = ?", (profile_id,)
            ).fetchone()[0],
        }


def validate(db_path: Path, template_path: Path) -> dict[str, Any]:
    expected = read_template(template_path)
    report: dict[str, Any] = {"db": str(db_path), "template": str(template_path), "checks": []}

    def check(name: str, passed: bool, detail: Any) -> None:
        report["checks"].append({"name": name, "passed": bool(passed), "detail": detail})

    with connect(db_path) as connection:
        tables = {
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'"
            )
        }
        required_tables = {
            "schema_meta",
            "cas_s2_profile",
            "cas_s2_source",
            "cas_s2_condition",
            "cas_s2_classification",
            "cas_s2_label_element",
            "cas_s2_hazard_summary",
            "cas_s2_fact",
            "cas_s2_review",
            "cas_s2_skeleton_field",
            "cas_s2_result_field",
        }
        check("required_tables", required_tables <= tables, sorted(required_tables - tables))

        fk_errors = [tuple(row) for row in connection.execute("PRAGMA foreign_key_check")]
        check("foreign_key_check", not fk_errors, fk_errors)

        profile = connection.execute(
            "SELECT * FROM cas_s2_profile WHERE cas_no = ? AND is_current = 1", (CAS_EXPECTED,)
        ).fetchone()
        check("one_current_nmp_profile", profile is not None, dict(profile) if profile else None)
        if profile is None:
            report["passed"] = False
            return report

        profile_id = int(profile["profile_id"])
        standard_rows = connection.execute(
            """
            SELECT field_key, field_label_zh, value_text, value_status, display_order
            FROM v_cas_s2_current_standard_result
            WHERE profile_id = ? ORDER BY display_order
            """,
            (profile_id,),
        ).fetchall()
        expected_field_keys = [key for key, _ in STANDARD_SKELETON_FIELDS]
        actual_field_keys = [row["field_key"] for row in standard_rows]
        check("fixed_standard_skeleton_13_fields", actual_field_keys == expected_field_keys, {
            "actual_count": len(actual_field_keys), "actual": actual_field_keys, "expected": expected_field_keys,
        })
        check("standard_field_values_are_explicit", all(str(row["value_text"] or "").strip() for row in standard_rows), [dict(row) for row in standard_rows])
        result_text_errors = []
        for row in standard_rows:
            try:
                validate_standard_result_text(row["field_key"], row["value_text"])
            except ValueError as exc:
                result_text_errors.append(str(exc))
        check("standard_fields_contain_direct_results_only", not result_text_errors, result_text_errors)
        summary_values = {
            row["field_key"]: row["value_text"]
            for row in standard_rows
            if row["field_key"] in {"physical_chemical_hazard", "health_hazard", "environmental_hazard", "other_hazard"}
        }
        check("blank_template_summaries_are_no_data", all(value == "无数据" for value in summary_values.values()), summary_values)
        checks = {
            "cas_no": profile["cas_no"] == expected["identity"]["cas_no"],
            "standard_name_zh": profile["standard_name_zh"] == expected["identity"]["standard_name_zh"],
            "ec_no": profile["ec_no"] == expected["identity"]["ec_no"],
            "index_no": profile["index_no"] == expected["identity"]["index_no"],
        }
        check("identity_matches_template", all(checks.values()), checks)
        profile_state_ok = (
            profile["profile_version"] == "1.1.0"
            and profile["profile_status"] == "manual_review"
            and "GB/T 17519-2013" in profile["regulatory_version"]
            and "EU CLP 2016/1179" in profile["regulatory_version"]
        )
        check("nmp_regulatory_profile_versioned", profile_state_ok, {
            "profile_version": profile["profile_version"],
            "profile_status": profile["profile_status"],
            "regulatory_version": profile["regulatory_version"],
        })

        classifications = [
            (row["classification_text"], row["h_code"])
            for row in connection.execute(
                "SELECT classification_text, h_code FROM cas_s2_classification WHERE profile_id = ? ORDER BY display_order",
                (profile_id,),
            )
        ]
        expected_classifications = []
        for line in expected["classification_text"].splitlines():
            if line.strip():
                _, classification_text, _, h_code = classification_parts(line)
                expected_classifications.append((classification_text, h_code))
        check("classifications_match_template", classifications == expected_classifications, {"actual": classifications, "expected": expected_classifications})

        pictograms = [
            row["element_code"]
            for row in connection.execute(
                "SELECT element_code FROM cas_s2_label_element WHERE profile_id = ? AND element_type = 'pictogram' ORDER BY display_order",
                (profile_id,),
            )
        ]
        expected_pictograms = split_codes(expected["pictogram_text"])
        check("pictograms_match_template", pictograms == expected_pictograms, {"actual": pictograms, "expected": expected_pictograms})

        signal = connection.execute(
            "SELECT element_text_zh FROM cas_s2_label_element WHERE profile_id = ? AND element_type = 'signal_word'",
            (profile_id,),
        ).fetchone()
        check("signal_word_matches_template", (signal[0] if signal else "") == expected["signal_word"], {"actual": signal[0] if signal else "", "expected": expected["signal_word"]})

        h_rows = [
            (row["element_code"], row["element_text_zh"])
            for row in connection.execute(
                "SELECT element_code, element_text_zh FROM cas_s2_label_element WHERE profile_id = ? AND element_type = 'hazard_statement' ORDER BY display_order",
                (profile_id,),
            )
        ]
        p_rows = [
            (row["element_code"], row["element_text_zh"])
            for row in connection.execute(
                "SELECT element_code, element_text_zh FROM cas_s2_label_element WHERE profile_id = ? AND element_type = 'precautionary_statement' ORDER BY display_order",
                (profile_id,),
            )
        ]
        check("hazard_statements_match_template", h_rows == [(a, b) for a, b, _ in parse_code_lines(expected["hazard_statement_text"], H_TEXT_RE)], {"actual": h_rows, "expected": [(a, b) for a, b, _ in parse_code_lines(expected["hazard_statement_text"], H_TEXT_RE)]})
        check("precautionary_statements_match_template", p_rows == [(a, b) for a, b, _ in parse_code_lines(expected["precautionary_statement_text"], P_TEXT_RE)], {"actual": p_rows, "expected": [(a, b) for a, b, _ in parse_code_lines(expected["precautionary_statement_text"], P_TEXT_RE)]})

        summary_rows = {
            row["hazard_domain"]: (row["summary_text"], row["value_status"])
            for row in connection.execute(
                "SELECT hazard_domain, summary_text, value_status FROM cas_s2_hazard_summary WHERE profile_id = ?",
                (profile_id,),
            )
        }
        summaries_ok = all(summary and status == "derived_candidate" for summary, status in summary_rows.values())
        check(
            "enriched_summaries_are_traceable_candidates",
            summaries_ok,
            {domain: {"has_text": bool(summary), "value_status": status} for domain, (summary, status) in summary_rows.items()},
        )

        condition = connection.execute(
            "SELECT condition_type, status, raw_text FROM cas_s2_condition WHERE profile_id = ? AND condition_key = 'profile_default'",
            (profile_id,),
        ).fetchone()
        condition_ok = bool(condition and condition["condition_type"] == "not_provided" and condition["status"] == "unknown")
        check("missing_threshold_is_not_executable", condition_ok, dict(condition) if condition else None)

        intrinsic = connection.execute(
            "SELECT condition_id, status, scope FROM cas_s2_condition WHERE profile_id = ? AND condition_key = 'substance_intrinsic'",
            (profile_id,),
        ).fetchone()
        intrinsic_ok = bool(intrinsic and intrinsic["status"] == "confirmed" and intrinsic["scope"] == "traceable_substance_classification")
        check("intrinsic_classification_condition_is_explicit", intrinsic_ok, dict(intrinsic) if intrinsic else None)

        thresholds = {
            row["condition_key"]: (row["condition_type"], row["lower_operator"], row["lower_value"], row["unit"], row["scope"], row["status"])
            for row in connection.execute(
                "SELECT condition_key, condition_type, lower_operator, lower_value, unit, scope, status FROM cas_s2_condition WHERE profile_id = ?",
                (profile_id,),
            )
        }
        expected_thresholds = {
            "cn_disclosure_reproductive_ge_0_1": ("lower_bound", ">=", 0.1, "%", "CN_mixture_component_disclosure", "confirmed"),
            "cn_disclosure_skin_ge_1": ("lower_bound", ">=", 1.0, "%", "CN_mixture_component_disclosure", "confirmed"),
            "cn_disclosure_eye_ge_1": ("lower_bound", ">=", 1.0, "%", "CN_mixture_component_disclosure", "confirmed"),
            "cn_disclosure_stot_se_ge_1": ("lower_bound", ">=", 1.0, "%", "CN_mixture_component_disclosure", "confirmed"),
            "eu_clp_nmp_stot_se_ge_10": ("lower_bound", ">=", 10.0, "%", "EU_CLP_mixture_classification", "confirmed"),
            "eu_reach_nmp_ge_0_3_restriction": ("lower_bound", ">=", 0.3, "%", "EU_REACH_regulatory_restriction", "confirmed"),
        }
        threshold_ok = all(thresholds.get(key) == value for key, value in expected_thresholds.items())
        historical_ok = thresholds.get("eu_clp_nmp_repr_ge_5_historical") == ("manual_review", ">=", 5.0, "%", "EU_CLP_historical_superseded", "manual_review")
        check("active_thresholds_have_scope_and_version", threshold_ok, {"actual": thresholds, "expected": expected_thresholds})
        check("historical_threshold_is_not_executable", historical_ok, thresholds.get("eu_clp_nmp_repr_ge_5_historical"))

        fact_count = connection.execute(
            "SELECT COUNT(*) FROM cas_s2_fact WHERE profile_id = ?", (profile_id,)
        ).fetchone()[0]
        not_listed = connection.execute(
            "SELECT value_status FROM cas_s2_fact WHERE profile_id = ? AND fact_code = 'cn_gbz2_1_nmp_listing'",
            (profile_id,),
        ).fetchone()
        check("structured_fact_library_populated", fact_count >= 30, {"fact_count": fact_count})
        check("gbz2_1_not_listed_is_explicit", bool(not_listed and not_listed[0] == "not_listed"), dict(not_listed) if not_listed else None)

        source_count = connection.execute(
            "SELECT source_count FROM v_cas_s2_current_profile WHERE cas_no = ?", (CAS_EXPECTED,)
        ).fetchone()[0]
        distinct_source_count = connection.execute(
            "SELECT COUNT(DISTINCT source_file_path || '|' || source_hash_sha256) FROM cas_s2_source WHERE profile_id = ? AND source_file_path <> '' AND source_file_path NOT LIKE 'remote://%'",
            (profile_id,),
        ).fetchone()[0]
        check("source_count_is_distinct_file_count", source_count == distinct_source_count == 11, {"view_count": source_count, "distinct_local_file_count": distinct_source_count})

    report["passed"] = all(item["passed"] for item in report["checks"])
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="CAS Section 2 结果库骨架工具")
    subparsers = parser.add_subparsers(dest="command", required=True)

    init_parser = subparsers.add_parser("init", help="初始化空数据库骨架")
    init_parser.add_argument("--db", type=Path, default=DEFAULT_DB)

    reset_parser = subparsers.add_parser("reset", help="在已有备份后重置本结果库")
    reset_parser.add_argument("--db", type=Path, default=DEFAULT_DB)

    seed_parser = subparsers.add_parser("seed-nmp", help="导入 NMP Excel 示例")
    seed_parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    seed_parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)

    enrich_parser = subparsers.add_parser("enrich-nmp", help="补充 NMP 法规来源、结构化事实和含量条件")
    enrich_parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    enrich_parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    enrich_parser.add_argument("--report", type=Path)

    nep_parser = subparsers.add_parser("enrich-nep", help="按已确认 NMP 骨架补充 NEP CAS 结果档案")
    nep_parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    nep_parser.add_argument(
        "--identity-db",
        type=Path,
        default=Path(r"F:\正式项目与模块化内容\冠志\MSDS\03-数据库\正式库\Data Base\cas_library.db"),
    )
    nep_parser.add_argument("--report", type=Path)

    validate_nep_parser = subparsers.add_parser("validate-nep", help="校验 NEP CAS 结果档案")
    validate_nep_parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    validate_nep_parser.add_argument(
        "--identity-db",
        type=Path,
        default=Path(r"F:\正式项目与模块化内容\冠志\MSDS\03-数据库\正式库\Data Base\cas_library.db"),
    )
    validate_nep_parser.add_argument("--report", type=Path)

    validate_parser = subparsers.add_parser("validate", help="校验 NMP 示例与 Excel 一致性")
    validate_parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    validate_parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    validate_parser.add_argument("--report", type=Path)

    args = parser.parse_args()
    if args.command == "init":
        initialize(args.db)
        print(json.dumps({"command": "init", "db": str(args.db), "passed": True}, ensure_ascii=False, indent=2))
        return 0
    if args.command == "reset":
        reset_database(args.db)
        print(json.dumps({"command": "reset", "db": str(args.db), "passed": True}, ensure_ascii=False, indent=2))
        return 0
    if args.command == "seed-nmp":
        result = insert_nmp(args.db, args.template)
        print(json.dumps({"command": "seed-nmp", "passed": True, **result}, ensure_ascii=False, indent=2))
        return 0
    if args.command == "enrich-nmp":
        from cas_section2_result_enrichment import enrich_nmp

        result = enrich_nmp(args.db, args.template)
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"command": "enrich-nmp", "passed": True, **result}, ensure_ascii=False, indent=2))
        return 0

    if args.command == "enrich-nep":
        from cas_section2_result_nep import enrich_nep

        result = enrich_nep(args.db, args.identity_db)
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"command": "enrich-nep", "passed": True, **result}, ensure_ascii=False, indent=2))
        return 0

    if args.command == "validate-nep":
        from cas_section2_result_nep import validate_nep

        report = validate_nep(args.db, args.identity_db)
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report.get("passed") else 1

    report = validate(args.db, args.template)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report.get("passed") else 1


if __name__ == "__main__":
    raise SystemExit(main())
